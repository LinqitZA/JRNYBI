import io

import openpyxl

from tests import BaseTestCase

URL = "/api/jrny/excel-export"
GRID = {
    "title": "Shoprite Weekly Invoice Count",
    "header_rows": 1,
    "rows": [
        [{"v": "CustomerName", "t": "s", "s": "header"}, {"v": "2026-40", "t": "s", "s": "header"}],
        [{"v": "Shoprite", "t": "s", "s": "rowhead"}, {"v": 5, "t": "n", "z": "#,##0"}],
    ],
}


class TestExcelExport(BaseTestCase):
    def test_returns_the_workbook_as_a_download(self):
        rv = self.make_request("post", URL, data=GRID, user=self.factory.user)
        self.assertEqual(rv.status_code, 200)
        self.assertEqual(rv.mimetype, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        self.assertIn('filename="Shoprite Weekly Invoice Count.xlsx"', rv.headers["Content-Disposition"])
        ws = openpyxl.load_workbook(io.BytesIO(rv.data)).active
        self.assertEqual(ws["A1"].value, "Shoprite Weekly Invoice Count")
        self.assertEqual(ws["B4"].value, 5)

    def test_non_ascii_title_gets_an_encoded_filename(self):
        rv = self.make_request("post", URL, data={**GRID, "title": "Orders – Capitec"}, user=self.factory.user)
        self.assertEqual(rv.status_code, 200)
        self.assertIn("filename*=UTF-8''Orders%20%E2%80%93%20Capitec.xlsx", rv.headers["Content-Disposition"])

    def test_invalid_grid_is_a_400_with_the_reason(self):
        rv = self.make_request("post", URL, data={"rows": [[{"v": "x", "t": "n"}]]}, user=self.factory.user)
        self.assertEqual(rv.status_code, 400)
        self.assertIn("number", rv.json["message"])

    def test_requires_login(self):
        rv = self.make_request("post", URL, data=GRID, user=False)
        self.assertIn(rv.status_code, (302, 401, 403))
