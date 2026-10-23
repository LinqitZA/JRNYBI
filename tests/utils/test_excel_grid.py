import datetime
import io
from unittest import TestCase

import openpyxl

from redash.utils.excel_grid import (
    GridError,
    build_workbook,
    safe_filename,
    validate_grid,
)


def cell(v, t="s", s=None, z=None):
    c = {"v": v, "t": t}
    if s:
        c["s"] = s
    if z:
        c["z"] = z
    return c


# The Shoprite Weekly Invoice Count shape: customers down, year-weeks across, totals.
PIVOT = {
    "title": "Shoprite Weekly Invoice Count",
    "subtitle": "Exported 21 Oct 2026 20:40",
    "sheet_name": "Report",
    "header_rows": 2,
    "header_cols": 1,
    "rows": [
        [None, cell("YearWeek", s="header"), cell("2026-40", s="header"), cell("2026-41", s="header"),
         cell("Totals", s="header")],
        [cell("CustomerName", s="header"), None, None, None, None],
        [cell("Checkers", s="rowhead"), None, cell(3, "n", "body", "#,##0"), None, cell(3, "n", "total", "#,##0")],
        [cell("Shoprite", s="rowhead"), None, cell(5, "n", "body", "#,##0"), cell(7, "n", "body", "#,##0"),
         cell(12, "n", "total", "#,##0")],
        [cell("Totals", s="total"), None, cell(8, "n", "total", "#,##0"), cell(7, "n", "total", "#,##0"),
         cell(15, "n", "total", "#,##0")],
    ],
    "merges": [[0, 4, 1, 4], [2, 0, 2, 1], [3, 0, 3, 1], [4, 0, 4, 1]],
}


def read(grid):
    return openpyxl.load_workbook(io.BytesIO(build_workbook(validate_grid(grid))))


class TestBuildWorkbook(TestCase):
    def test_title_subtitle_then_grid_after_a_blank_row(self):
        ws = read(PIVOT)["Report"]
        self.assertEqual(ws["A1"].value, "Shoprite Weekly Invoice Count")
        self.assertTrue(ws["A1"].font.b)
        self.assertEqual(ws["A2"].value, "Exported 21 Oct 2026 20:40")
        self.assertIsNone(ws["A3"].value)
        self.assertEqual(ws["B4"].value, "YearWeek")  # grid row 0 starts at Excel row 4
        self.assertEqual(ws["A5"].value, "CustomerName")

    def test_values_are_numbers_with_their_format_and_totals_are_bold(self):
        ws = read(PIVOT)["Report"]
        self.assertEqual(ws["C7"].value, 5)
        self.assertIsInstance(ws["C7"].value, int)
        self.assertEqual(ws["C7"].number_format, "#,##0")
        self.assertFalse(ws["C7"].font.b)
        self.assertEqual(ws["E8"].value, 15)
        self.assertTrue(ws["E8"].font.b)
        self.assertTrue(ws["A6"].font.b)  # row header
        self.assertIsNone(ws["D6"].value)  # empty pivot cell stays empty

    def test_merges_freeze_and_borders(self):
        ws = read(PIVOT)["Report"]
        merged = {str(r) for r in ws.merged_cells.ranges}
        self.assertEqual(merged, {"E4:E5", "A6:B6", "A7:B7", "A8:B8"})
        self.assertEqual(ws.freeze_panes, "B6")  # below 2 header rows, right of 1 header column
        self.assertEqual(ws["C7"].border.left.style, "thin")
        self.assertEqual(ws["D6"].border.left.style, "thin")  # empty grid cells are bordered too

    def test_column_widths_follow_the_text(self):
        ws = read(PIVOT)["Report"]
        self.assertGreaterEqual(ws.column_dimensions["A"].width, len("CustomerName"))
        self.assertLessEqual(ws.column_dimensions["A"].width, 60)

    def test_a_label_merged_across_columns_gets_room_for_its_text(self):
        long_name = "Checkers Hyper Gauteng Division"
        grid = {
            "header_rows": 1,
            "header_cols": 2,
            "rows": [
                [cell("Customer", s="header"), None, cell("2026-40", s="header")],
                [cell(long_name, s="rowhead"), None, cell(3, "n")],
            ],
            "merges": [[1, 0, 1, 1]],
        }
        ws = read(grid).active
        span = ws.column_dimensions["A"].width + ws.column_dimensions["B"].width
        self.assertGreaterEqual(span, len(long_name))

    def test_empty_cells_in_the_header_corner_look_like_headers(self):
        ws = read(PIVOT)["Report"]
        self.assertEqual(ws["A4"].fill.fgColor.rgb, ws["B4"].fill.fgColor.rgb)  # corner = header fill
        self.assertNotEqual(ws["D6"].fill.fgColor.rgb, ws["B4"].fill.fgColor.rgb)  # empty body cell is not

    def test_text_that_looks_like_a_formula_stays_text(self):
        grid = {"rows": [[cell("=1+1")], [cell("@SUM(A1)")]], "header_rows": 0}
        ws = read(grid).active
        self.assertEqual(ws["A1"].value, "=1+1")
        self.assertEqual(ws["A1"].data_type, "s")
        self.assertEqual(ws["A2"].value, "@SUM(A1)")

    def test_dates_booleans_and_no_title(self):
        grid = {
            "header_rows": 1,
            "rows": [
                [cell("Date", s="header"), cell("When", s="header"), cell("Paid", s="header")],
                [cell("2026-10-21", "d", z="dd/mm/yyyy"), cell("2026-10-21T13:45:00", "d", z="dd/mm/yyyy hh:mm"),
                 cell(True, "b")],
            ],
        }
        ws = read(grid).active
        self.assertEqual(ws["A1"].value, "Date")  # no title: the grid starts in row 1
        self.assertEqual(ws["A2"].value, datetime.datetime(2026, 10, 21))
        self.assertEqual(ws["A2"].number_format, "dd/mm/yyyy")
        self.assertEqual(ws["B2"].value, datetime.datetime(2026, 10, 21, 13, 45))
        self.assertIs(ws["C2"].value, True)
        self.assertEqual(ws.freeze_panes, "A2")

    def test_sheet_name_is_cleaned(self):
        wb = read({"rows": [[cell("x")]], "sheet_name": "Orders: [Capitec] / 2026?" + "x" * 40})
        name = wb.sheetnames[0]
        self.assertLessEqual(len(name), 31)
        self.assertFalse(set(name) & set("[]:*?/\\"))


# The Income Statement (Drill Down) shape: group > description > code, a total outside the outline.
OUTLINE = {
    "header_rows": 1,
    "rows": [[cell("Group", s="header")], [cell("Sales")], [cell("Desc A")], [cell("100010")], [cell("100020")],
             [cell("Desc B")], [cell("100030")], [cell("Gross Income")]],
    "row_outline": {"levels": [0, 0, 1, 2, 2, 1, 2, 0], "collapsed": True},
}


class TestRowOutline(TestCase):
    def test_rows_are_grouped_and_open_collapsed_with_the_summary_row_above(self):
        ws = read(OUTLINE).worksheets[0]
        dims = ws.row_dimensions
        # Excel rows 1..8 (no title)
        self.assertEqual([dims[r].outline_level for r in range(1, 9)], [0, 0, 1, 2, 2, 1, 2, 0])
        self.assertEqual([bool(dims[r].hidden) for r in range(1, 9)],
                         [False, False, True, True, True, True, True, False])
        self.assertFalse(ws.sheet_properties.outlinePr.summaryBelow)
        self.assertEqual(ws["A2"].value, "Sales")  # every row is still in the file
        self.assertEqual(ws["A7"].value, "100030")

    def test_expanded_outline_and_title_offset(self):
        grid = dict(OUTLINE, title="Income Statement", row_outline={"levels": OUTLINE["row_outline"]["levels"],
                                                                     "collapsed": False})
        ws = read(grid).worksheets[0]
        dims = ws.row_dimensions
        # the grid starts at Excel row 3 (title + blank row)
        self.assertEqual([dims[r].outline_level for r in range(3, 11)], [0, 0, 1, 2, 2, 1, 2, 0])
        self.assertFalse(any(dims[r].hidden for r in range(3, 11)))

    def test_no_outline_by_default(self):
        ws = read(PIVOT)["Report"]
        self.assertTrue(all(not d.outline_level for d in ws.row_dimensions.values()))

    def test_rejects_a_bad_outline(self):
        rows = OUTLINE["rows"]
        for bad, text in (
            ("x", "row_outline"),
            ({"levels": [0, 1]}, "one level per row"),
            ({"levels": [0, 0, 1, 2, 2, 1, 9, 0]}, "0..7"),
            ({"levels": [0, 0, 1, 2, 2, 1, "2", 0]}, "0..7"),
            ({"levels": [0, 0, 1, 2, 2, 1, 2, 0], "collapsed": "yes"}, "collapsed"),
        ):
            with self.assertRaises(GridError) as ctx:
                validate_grid({"rows": rows, "row_outline": bad})
            self.assertIn(text, str(ctx.exception))


class TestValidateGrid(TestCase):
    def assertInvalid(self, grid, text):
        with self.assertRaises(GridError) as ctx:
            validate_grid(grid)
        self.assertIn(text, str(ctx.exception))

    def test_rejects_bad_shapes(self):
        self.assertInvalid([], "object")
        self.assertInvalid({"rows": "x"}, "rows")
        self.assertInvalid({"rows": [["x"]]}, "cell")
        self.assertInvalid({"rows": [[{"v": 1, "t": "q"}]]}, "type")
        self.assertInvalid({"rows": [[{"v": "abc", "t": "n"}]]}, "number")
        self.assertInvalid({"rows": [[{"v": "21/10/2026", "t": "d"}]]}, "date")
        self.assertInvalid({"rows": [[cell("x", s="loud")]]}, "style")

    def test_rejects_bad_merges(self):
        rows = [[cell("a"), None], [None, None]]
        self.assertInvalid({"rows": rows, "merges": [[0, 0, 5, 0]]}, "merge")
        self.assertInvalid({"rows": rows, "merges": [[0, 0, 1, 1], [1, 1, 1, 1]]}, "overlap")
        self.assertInvalid({"rows": rows, "merges": [[1, 0, 0, 0]]}, "merge")

    def test_size_limits(self):
        self.assertInvalid({"rows": [[None] * 16385]}, "columns")
        self.assertInvalid({"rows": [[None]], "header_rows": 201}, "header")


class TestSafeFilename(TestCase):
    def test_keeps_a_readable_name(self):
        self.assertEqual(safe_filename("Shoprite Weekly Invoice Count"), "Shoprite Weekly Invoice Count.xlsx")
        self.assertEqual(safe_filename('a/b\\c:"d"\n'), "a b c d.xlsx")
        self.assertEqual(safe_filename(""), "report.xlsx")
